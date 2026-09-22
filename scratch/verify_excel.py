import sys
import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook("Lista_Inmobiliarias_y_Telefonos_ICDE.xlsx")
print("Sheet names:", wb.sheetnames)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\n--- SHEET: {sheetname} ---")
    data = list(ws.iter_rows(values_only=True))
    print(f"Total rows: {len(data)}")
    for row in data[:15]:
        print(row)
