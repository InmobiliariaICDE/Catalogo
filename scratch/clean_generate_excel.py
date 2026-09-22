import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

sys.stdout.reconfigure(encoding='utf-8')

# Find Matriz file
matriz_file = [f for f in os.listdir('.') if f.startswith('Matriz') and f.endswith('.xlsx')][0]
print(f"Reading file: {matriz_file}")

wb_in = openpyxl.load_workbook(matriz_file, data_only=True, read_only=True)
ws_in = wb_in['Base de Datos']

rows = list(ws_in.iter_rows(values_only=True))
header = rows[0]

idx_cod = header.index('Código') if 'Código' in header else 0
idx_nombre = header.index('Nombre') if 'Nombre' in header else 1
idx_tipo = header.index('Tipo de inmueble') if 'Tipo de inmueble' in header else 2
idx_barrio = header.index('Barrio') if 'Barrio' in header else 3
idx_inmo = header.index('Inmobiliaria')
idx_cel1 = header.index('CELULAR')
idx_cel2 = header.index('CELULAR 2')
idx_prop = header.index('PROPIETARIO') if 'PROPIETARIO' in header else -1

def clean_phone(val):
    if val is None:
        return ""
    val_str = str(val).strip()
    if val_str.endswith('.0'):
        val_str = val_str[:-2]
    if val_str in ["0", "nan", "None", "0.0", "?????"]:
        return ""
    return val_str

detail_rows = []

for idx, r in enumerate(rows[1:], start=2):
    inmo_raw = str(r[idx_inmo]).strip() if r[idx_inmo] is not None else ""
    cel1 = clean_phone(r[idx_cel1])
    cel2 = clean_phone(r[idx_cel2])
    prop = str(r[idx_prop]).strip() if idx_prop != -1 and r[idx_prop] is not None else ""
    cod = str(r[idx_cod]).strip() if r[idx_cod] is not None else ""
    nombre = str(r[idx_nombre]).strip() if r[idx_nombre] is not None else ""
    tipo = str(r[idx_tipo]).strip() if r[idx_tipo] is not None else ""
    barrio = str(r[idx_barrio]).strip() if r[idx_barrio] is not None else ""
    
    # Exclude rows where Inmobiliaria is blank or None
    if inmo_raw and inmo_raw.lower() != "none":
        detail_rows.append({
            'Fila_Matriz': idx,
            'Código': cod,
            'Nombre_Inmueble': nombre,
            'Tipo_Inmueble': tipo,
            'Barrio': barrio,
            'Inmobiliaria': inmo_raw,
            'Celular_1': cel1,
            'Celular_2': cel2,
            'Propietario_Notas': prop
        })

df_detail = pd.DataFrame(detail_rows)

# Grouping for Summary Sheet
summary_map = {}

for idx, row in df_detail.iterrows():
    inmo = row['Inmobiliaria']
    c1 = row['Celular_1']
    c2 = row['Celular_2']
    cod = row['Código']
    
    if inmo not in summary_map:
        summary_map[inmo] = {
            'Inmobiliaria': inmo,
            'Cantidad_Inmuebles': 0,
            'Celulares_1_Set': set(),
            'Celulares_2_Set': set(),
            'Todos_Celulares_Set': set(),
            'Codigos_List': []
        }
    
    summary_map[inmo]['Cantidad_Inmuebles'] += 1
    if cod:
        summary_map[inmo]['Codigos_List'].append(cod)
    if c1:
        summary_map[inmo]['Celulares_1_Set'].add(c1)
        summary_map[inmo]['Todos_Celulares_Set'].add(c1)
    if c2:
        summary_map[inmo]['Celulares_2_Set'].add(c2)
        summary_map[inmo]['Todos_Celulares_Set'].add(c2)

summary_rows = []
for inmo, data in summary_map.items():
    cels1_str = " | ".join(sorted(list(data['Celulares_1_Set'])))
    cels2_str = " | ".join(sorted(list(data['Celulares_2_Set'])))
    all_cels_str = " | ".join(sorted(list(data['Todos_Celulares_Set'])))
    cods_str = ", ".join(data['Codigos_List'][:10])
    if len(data['Codigos_List']) > 10:
        cods_str += f" (+{len(data['Codigos_List']) - 10} más)"
        
    summary_rows.append({
        'Inmobiliaria / Contacto': inmo,
        'Cant. Inmuebles': data['Cantidad_Inmuebles'],
        'Celular(es) Principal(es)': cels1_str,
        'Celular(es) Secundario(s)': cels2_str,
        'Todos los Celulares Registrados': all_cels_str if all_cels_str else "Sin número registrado",
        'Muestra Códigos Inmuebles': cods_str
    })

df_summary = pd.DataFrame(summary_rows)
# Sort by Cant. Inmuebles descending
df_summary = df_summary.sort_values(by=['Cant. Inmuebles', 'Inmobiliaria / Contacto'], ascending=[False, True])

# Create Excel Workbook
output_filename = "Lista_Inmobiliarias_y_Telefonos_ICDE.xlsx"
wb_out = openpyxl.Workbook()

# Sheet 1: Resumen Inmobiliarias
ws1 = wb_out.active
ws1.title = "Resumen Inmobiliarias"

# Sheet 2: Detalle por Inmueble
ws2 = wb_out.create_sheet(title="Detalle Inmuebles")

# Styles
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
fill_header_blue = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
fill_header_navy = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
font_title = Font(name="Calibri", size=15, bold=True, color="1F4E78")
font_subtitle = Font(name="Calibri", size=10, italic=True, color="595959")
font_regular = Font(name="Calibri", size=11)

align_left = Alignment(horizontal="left", vertical="center")
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# --- WRITE SHEET 1 ---
ws1.cell(row=1, column=1, value="DIRECTORIO DE INMOBILIARIAS Y CONTACTOS - BASE DE DATOS ICDE").font = font_title
ws1.cell(row=2, column=1, value=f"Total: {len(df_summary)} Inmobiliarias/Contactos | {len(df_detail)} Inmuebles").font = font_subtitle

headers_s1 = ['Inmobiliaria / Contacto', 'Cant. Inmuebles', 'Celular(es) Principal(es)', 'Celular(es) Secundario(s)', 'Todos los Celulares Registrados', 'Muestra Códigos Inmuebles']

header_row_1 = 3
for col_idx, h_name in enumerate(headers_s1, start=1):
    cell = ws1.cell(row=header_row_1, column=col_idx, value=h_name)
    cell.font = font_header
    cell.fill = fill_header_blue
    cell.alignment = align_center

for row_idx, data_dict in enumerate(df_summary.to_dict('records'), start=4):
    row_vals = [
        data_dict['Inmobiliaria / Contacto'],
        data_dict['Cant. Inmuebles'],
        data_dict['Celular(es) Principal(es)'],
        data_dict['Celular(es) Secundario(s)'],
        data_dict['Todos los Celulares Registrados'],
        data_dict['Muestra Códigos Inmuebles']
    ]
    for col_idx, val in enumerate(row_vals, start=1):
        cell = ws1.cell(row=row_idx, column=col_idx, value=val)
        cell.font = font_regular
        cell.border = thin_border
        cell.alignment = align_center if col_idx == 2 else align_left

# --- WRITE SHEET 2 ---
ws2.cell(row=1, column=1, value="DETALLE DE INMUEBLES E INMOBILIARIAS ASOCIADAS").font = font_title
ws2.cell(row=2, column=1, value="Detalle fila por fila de la Base de Datos ICDE con Inmobiliaria y Celulares asociados").font = font_subtitle

headers_s2 = ['Fila Matriz', 'Código', 'Nombre Inmueble', 'Tipo Inmueble', 'Barrio', 'Inmobiliaria', 'Celular 1', 'Celular 2', 'Propietario / Contacto']

header_row_2 = 3
for col_idx, h_name in enumerate(headers_s2, start=1):
    cell = ws2.cell(row=header_row_2, column=col_idx, value=h_name)
    cell.font = font_header
    cell.fill = fill_header_navy
    cell.alignment = align_center

for row_idx, data_dict in enumerate(df_detail.to_dict('records'), start=4):
    row_vals = [
        data_dict['Fila_Matriz'],
        data_dict['Código'],
        data_dict['Nombre_Inmueble'],
        data_dict['Tipo_Inmueble'],
        data_dict['Barrio'],
        data_dict['Inmobiliaria'],
        data_dict['Celular_1'],
        data_dict['Celular_2'],
        data_dict['Propietario_Notas']
    ]
    for col_idx, val in enumerate(row_vals, start=1):
        cell = ws2.cell(row=row_idx, column=col_idx, value=val)
        cell.font = font_regular
        cell.border = thin_border
        cell.alignment = align_center if col_idx in [1, 2] else align_left

# Adjust column dimensions
ws1.column_dimensions['A'].width = 45
ws1.column_dimensions['B'].width = 18
ws1.column_dimensions['C'].width = 30
ws1.column_dimensions['D'].width = 30
ws1.column_dimensions['E'].width = 40
ws1.column_dimensions['F'].width = 35

ws2.column_dimensions['A'].width = 14
ws2.column_dimensions['B'].width = 14
ws2.column_dimensions['C'].width = 45
ws2.column_dimensions['D'].width = 25
ws2.column_dimensions['E'].width = 25
ws2.column_dimensions['F'].width = 35
ws2.column_dimensions['G'].width = 20
ws2.column_dimensions['H'].width = 20
ws2.column_dimensions['I'].width = 35

for ws in [ws1, ws2]:
    ws.views.sheetView[0].showGridLines = True

wb_out.save(output_filename)
print(f"Clean Excel successfully saved to: {os.path.abspath(output_filename)}")
