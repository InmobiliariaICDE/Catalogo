import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

file_path = "Lista_Inmobiliarias_y_Telefonos_ICDE.xlsx"
wb = openpyxl.load_workbook(file_path)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    rows = list(ws.iter_rows(values_only=True))
    print(f"Sheet: {sheetname} | Total rows: {len(rows)}")
    empty_rows = []
    for r_idx, r in enumerate(rows, start=1):
        if not r or all(v is None or str(v).strip() == "" for v in r):
            empty_rows.append(r_idx)
    print(f"Empty row count in {sheetname}: {len(empty_rows)}, Row indices: {empty_rows[:20]}")
