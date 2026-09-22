import os
import sys
import glob
import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding='utf-8')


# Find file matching Matriz*.xlsx
files = [f for f in os.listdir('.') if f.startswith('Matriz') and f.endswith('.xlsx')]
print(f"Found {len(files)} files:")
for f in files:
    print(f"File: {repr(f)}")

if not files:
    print("No Matriz file found!")
    exit(1)

target_file = files[0]

wb = openpyxl.load_workbook(target_file, data_only=True, read_only=True)
print("Sheet names:", wb.sheetnames)

# Find sheet matching 'base de datos' or similar (case insensitive)
sheet_name = None
for s in wb.sheetnames:
    if 'base' in s.lower() and 'dato' in s.lower():
        sheet_name = s
        break

if not sheet_name:
    print("Could not find base de datos sheet!")
    sheet_name = wb.sheetnames[0]

print(f"Using sheet: {repr(sheet_name)}")

ws = wb[sheet_name]
rows = list(ws.iter_rows(values_only=True))

if not rows:
    print("Sheet is empty!")
    exit(0)

print(f"Total rows: {len(rows)}")

# Print first few rows to inspect headers
for idx, r in enumerate(rows[:10]):
    # filter non-None values for readable preview
    non_empty = [(i, val) for i, val in enumerate(r) if val is not None]
    print(f"Row {idx}: {non_empty[:15]}")
